"""Base central de CMV: leitura do arquivo do Grupo e aplicação às plataformas, sem rede."""
import io
import json
import os
import tempfile
import unittest
from unittest import mock

import openpyxl

import cmv_base


def planilha(abas):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome, linhas in abas.items():
        ws = wb.create_sheet(nome)
        for r, l in enumerate(linhas, start=1):
            for c, v in enumerate(l, start=1):
                ws.cell(row=r, column=c, value=v)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


ARQ = planilha({
    "Todos": [["CMV Confirmado"], ["#", "SKU", "Descrição", "Tipo", "CMV (R$)", "Status"],
              [1, "KIT-A", "Kit A", "kit", 49, "ok"], [2, "kit-b", "Kit B", "kit", "11,60", "ok"], [3, "SEM-CUSTO", "x", "kit", None, ""]],
    "Clones": [["Grupos"], ["Produto", "SKU Clone", "CMV (R$)", "Obs"], ["▸ Torneira", None, None, None],
               [None, "TORB-FBA", 13.5, "ok"], [None, "KIT-A", 50, "conflito"]],
})


class Leitura(unittest.TestCase):
    def test_le_todas_as_abas_e_trata_numero_com_virgula(self):
        r = cmv_base.ler_planilha(ARQ)
        self.assertEqual(r["itens"]["KIT-B"]["cmv"], 11.6)
        self.assertEqual(r["itens"]["TORB-FBA"]["cmv"], 13.5)
        self.assertEqual([a["skus_lidos"] for a in r["abas"]], [2, 1])

    def test_conflito_mantem_a_primeira_aba(self):
        r = cmv_base.ler_planilha(ARQ)
        self.assertEqual(r["itens"]["KIT-A"]["cmv"], 49.0)
        self.assertEqual(r["conflitos"][0]["cmv_ignorado"], 50.0)

    def test_cmv_vazio_e_ignorado_e_arquivo_sem_colunas_falha(self):
        self.assertEqual(cmv_base.ler_planilha(ARQ)["ignorados"][0]["sku"], "SEM-CUSTO")
        with self.assertRaises(Exception):
            cmv_base.importar(planilha({"x": [["a", "b"], [1, 2]]}))


class Aplicacao(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.TemporaryDirectory()
        self.addCleanup(self.d.cleanup)
        p = mock.patch.object(cmv_base, "PATH", os.path.join(self.d.name, "base.json"))
        p.start()
        self.addCleanup(p.stop)
        cmv_base.importar(ARQ, "teste.xlsx")

    def test_busca_sem_diferenciar_caixa(self):
        self.assertEqual(cmv_base.cmv_de("Kit-A"), 49.0)
        self.assertIsNone(cmv_base.cmv_de("nao-existe"))

    def test_ml_simula_e_so_grava_com_confirmar(self):
        pp = os.path.join(self.d.name, "produtos.json")
        json.dump({"MLB1": {"conta": "meli01", "sku": "kit-a", "cmv_unit": 10},
                   "MLB2": {"conta": "meli01", "sku": "TORB-FBA", "cmv_unit": 13.5},
                   "MLB3": {"conta": "meli01", "sku": "", "cmv_unit": 1},
                   "MLB4": {"conta": "meli02", "sku": "KIT-A", "cmv_unit": 5}}, open(pp, "w"))
        r = cmv_base.aplicar_ml("meli01", False, pp)
        self.assertEqual((r["mudariam"], r["ja_iguais"], r["sem_sku"], r["aplicado"]), (1, 1, 1, False))
        self.assertEqual(json.load(open(pp))["MLB1"]["cmv_unit"], 10)
        r = cmv_base.aplicar_ml("meli01", True, pp)
        d = json.load(open(pp))
        self.assertTrue(r["aplicado"])
        self.assertEqual((d["MLB1"]["cmv_unit"], d["MLB4"]["cmv_unit"]), (49.0, 5))   # outra conta intocada

    def test_amazon_exige_definir_st_para_sku_novo(self):
        cp = os.path.join(self.d.name, "custos.json")
        r = cmv_base.aplicar_amazon("amazon02", None, False, cp, (0.1072, 0.1487), ["KIT-A", "NAO-TEM"])
        self.assertEqual((r["sem_imposto_definido"], r["sem_cmv_na_base"], r["mudariam"]), (["KIT-A"], ["NAO-TEM"], 0))
        r = cmv_base.aplicar_amazon("amazon02", False, True, cp, (0.1072, 0.1487), ["KIT-A"])
        c = json.load(open(cp))["amazon02"]["KIT-A"]
        self.assertEqual((c["cmv_unit"], c["imposto_pct"], c["tem_st"]), (49.0, 0.1487, False))

    def test_amazon_preserva_imposto_ja_cadastrado(self):
        cp = os.path.join(self.d.name, "custos.json")
        json.dump({"amazon01": {"kit-a": {"cmv_unit": 1, "imposto_pct": 0.115, "tem_st": True}}}, open(cp, "w"))
        cmv_base.aplicar_amazon("amazon01", None, True, cp, (0.115, 0.115), ["KIT-A"])
        c = json.load(open(cp))["amazon01"]
        self.assertEqual(list(c), ["kit-a"])
        self.assertEqual((c["kit-a"]["cmv_unit"], c["kit-a"]["imposto_pct"]), (49.0, 0.115))


class AmazonUsaABase(unittest.TestCase):
    def test_custos_com_base_completa_so_o_que_falta(self):
        from marketplaces.amazon import rentabilidade
        d = tempfile.TemporaryDirectory()
        self.addCleanup(d.cleanup)
        with mock.patch.object(cmv_base, "PATH", os.path.join(d.name, "base.json")),              mock.patch.object(rentabilidade, "CUSTOS_PATH", os.path.join(d.name, "custos.json")),              mock.patch.object(cmv_base, "aliquotas_da_conta", return_value=(0.1072, 0.1487)):
            cmv_base.importar(ARQ, "t.xlsx")
            json.dump({"amazon02": {"KIT-A": {"cmv_unit": 1.0, "imposto_pct": 0.5}}}, open(os.path.join(d.name, "custos.json"), "w"))
            c = rentabilidade.custos_com_base("amazon02", ["kit-a", "KIT-B", "NAO-TEM"])
        self.assertEqual(c["kit-a"]["cmv_unit"], 1.0)            # o arquivo explícito tem prioridade
        self.assertEqual((c["KIT-B"]["cmv_unit"], c["KIT-B"]["imposto_pct"]), (11.6, 0.1072))
        self.assertNotIn("NAO-TEM", c)


if __name__ == "__main__":
    unittest.main()
