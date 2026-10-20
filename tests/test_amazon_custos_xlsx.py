"""Modelo e importação de CMV/imposto da Amazon: sem rede."""
import io
import json
import os
import tempfile
import unittest
from unittest import mock

import openpyxl

from marketplaces.amazon import api, custos_xlsx, rentabilidade

SKUS = [{"sku": "KIT-A", "asin": "B01", "titulo": "Kit A"}, {"sku": "kit-b", "asin": "B02", "titulo": "Kit B"}]


class Base(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        p = mock.patch.object(rentabilidade, "CUSTOS_PATH", os.path.join(self.dir.name, "custos.json"))
        p.start()
        self.addCleanup(p.stop)

    def planilha(self, linhas, com=10, sem=8):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "Alíquota COM ST", com, "SEM ST", sem
        for i, h in enumerate(custos_xlsx.HEADERS, start=1):
            ws.cell(row=4, column=i, value=h)
        for r, l in enumerate(linhas, start=5):
            for i, v in enumerate(l, start=1):
                ws.cell(row=r, column=i, value=v)
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()


class Modelo(Base):
    def test_modelo_tem_cabecalho_aliquotas_e_skus(self):
        wb = openpyxl.load_workbook(io.BytesIO(custos_xlsx.gerar_modelo("amazon01", SKUS)))
        ws = wb.active
        self.assertEqual([c.value for c in ws[4]][:5], custos_xlsx.HEADERS)
        self.assertEqual((ws["B2"].value, ws["D2"].value), (10.0, 8.0))
        self.assertEqual([ws["A5"].value, ws["A6"].value], ["KIT-A", "kit-b"])

    def test_modelo_traz_custos_atuais(self):
        with open(rentabilidade.CUSTOS_PATH, "w", encoding="utf-8") as f:
            json.dump({"amazon01": {"KIT-A": {"cmv_unit": 20.5, "imposto_pct": 0.10, "tem_st": True}}}, f)
        ws = openpyxl.load_workbook(io.BytesIO(custos_xlsx.gerar_modelo("amazon01", SKUS))).active
        self.assertEqual((ws["D5"].value, ws["E5"].value), (20.5, "COM ST"))


class Importar(Base):
    def test_ida_e_volta_grava_no_formato_da_rentabilidade(self):
        r = custos_xlsx.importar("amazon01", self.planilha([["KIT-A", "B01", "x", 20.5, "COM ST"], ["kit-b", "B02", "y", "R$ 8,25", "SEM ST"]]))
        self.assertEqual((r["novos"], r["atualizados"], r["nao_aceitos"]), (2, 0, []))
        c = rentabilidade.carregar_custos("amazon01")
        self.assertEqual(c["KIT-A"], {"cmv_unit": 20.5, "imposto_pct": 0.1, "tem_st": True})
        self.assertEqual(c["kit-b"]["cmv_unit"], 8.25)
        self.assertEqual(c["kit-b"]["imposto_pct"], 0.08)

    def test_reimportar_atualiza_e_nao_apaga_outros_skus_nem_contas(self):
        custos_xlsx.importar("amazon01", self.planilha([["KIT-A", "", "", 10, "COM ST"], ["OUTRO", "", "", 5, "SEM ST"]]))
        custos_xlsx.importar("amazon02", self.planilha([["Z", "", "", 3, "SEM ST"]]))
        r = custos_xlsx.importar("amazon01", self.planilha([["kit-a", "", "", 12, "COM ST"]]))
        self.assertEqual((r["atualizados"], r["novos"]), (1, 0))   # SKU em outra caixa atualiza o existente
        c = rentabilidade.carregar_custos("amazon01")
        self.assertEqual(c["KIT-A"]["cmv_unit"], 12.0)
        self.assertIn("OUTRO", c)
        self.assertIn("Z", rentabilidade.carregar_custos("amazon02"))

    def test_aliquotas_da_planilha_sobrepoem(self):
        custos_xlsx.importar("amazon01", self.planilha([["A", "", "", 1, "COM ST"], ["B", "", "", 1, "SEM ST"]], com=12, sem=6))
        c = rentabilidade.carregar_custos("amazon01")
        self.assertEqual((c["A"]["imposto_pct"], c["B"]["imposto_pct"]), (0.12, 0.06))

    def test_campos_obrigatorios_e_invalidos(self):
        r = custos_xlsx.importar("amazon01", self.planilha([["A", "", "", None, "COM ST"], ["B", "", "", 5, "TALVEZ"], ["C", "", "", -1, "SEM ST"], ["D", "", "", 5, "SEM ST"]]))
        self.assertEqual(r["novos"], 1)
        self.assertEqual(len(r["nao_aceitos"]), 3)

    def test_arquivo_invalido(self):
        with self.assertRaises(api.ErroAmazon):
            custos_xlsx.importar("amazon01", b"isto nao e xlsx")

    def test_sem_cabecalho(self):
        wb = openpyxl.Workbook()
        b = io.BytesIO()
        wb.save(b)
        with self.assertRaises(api.ErroAmazon):
            custos_xlsx.importar("amazon01", b.getvalue())

    def test_dre_usa_os_custos_importados(self):
        custos_xlsx.importar("amazon01", self.planilha([["S", "", "", 10, "COM ST"]]))
        agg = {"receita": 100.0, "receita_estornada": 0.0, "promocoes": 0.0, "promocoes_estornadas": 0.0, "comissao": -15.0,
               "comissao_estornada": 0.0, "taxa_fba": -10.0, "taxa_fba_estornada": 0.0, "outras_taxas": 0.0,
               "outras_taxas_estornadas": 0.0, "un_enviadas": 2, "un_devolvidas": 0}
        r = rentabilidade.margem_sku(agg, rentabilidade.carregar_custos("amazon01")["S"])
        self.assertTrue(r["completo"])
        self.assertEqual((r["imposto"], r["cmv"], r["margem"]), (-10.0, -20.0, 45.0))


if __name__ == "__main__":
    unittest.main()
