"""
reposicao_full.py — Protótipo de cobertura e sugestão de reposição do Full
MZCell Rentabilidade

SOMENTE LEITURA: só faz GET na API do Mercado Livre. Não altera o app.
Uso:  .\\venv\\Scripts\\python.exe reposicao_full.py [meli01,meli02,meli03]

Regras (ver memória do módulo / diag_full.py):
- Unidade de cálculo = inventory_id (vários MLBs podem dividir o mesmo estoque Full).
- Estoque considerado = available_quantity + transfer (em transferência entre CDs).
- Vendas = unidades dos pedidos (exclui invalid, cancelled e pack_splitted), fuso -04:00 como o parser.
- Janelas de 7, 15 e 30 dias terminando no último dia fechado (ontem).
- Demanda = 50% média 7d + 30% média 15d + 20% média 30d; em tendência de alta (7d/30d > 1,15) usa a maior média.
- Sugestão = max(0, ceil(META_DIAS × demanda − estoque considerado)).
"""

import asyncio
import math
import os
import sys
import time
import warnings
from collections import defaultdict
from datetime import datetime, timedelta, timezone

import httpx

import ml_api

warnings.filterwarnings("ignore")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
B = ml_api.ML_BASE_URL
ML_TZ = timezone(timedelta(hours=-4))  # mesmo fuso do parser_rentabilidade
META_DIAS = 20
PESOS = {7: 0.5, 15: 0.3, 30: 0.2}
ALTA, QUEDA = 1.15, 0.85
STATUS_FORA = {"invalid", "cancelled"}
SEM_MOVIMENTO = "sem_movimento"
STATUS_PT = {"active": "ativo", "paused": "pausado", "closed": "encerrado", "under_review": "em revisão"}


class ML:
    def __init__(self, conta: str):
        self.conta = conta
        self.cli = httpx.Client(timeout=30, verify=False)
        self.renovou = False

    def get(self, path: str, **params):
        espera, n429 = 2, 0
        while True:
            r = self.cli.get(B + path, params=params, headers=ml_api._headers(self.conta))
            if r.status_code == 401 and not self.renovou:
                self.renovou = True
                asyncio.run(ml_api.refresh_token(self.conta))
                continue
            if (r.status_code == 429 or r.status_code >= 500) and n429 < 6:
                n429 += 1
                time.sleep(espera)
                espera = min(espera * 2, 30)
                continue
            return r.status_code, (r.json() if r.headers.get("content-type", "").startswith("application/json") else None)


def detalhar(ml: ML, ids: list) -> dict:
    itens = {}
    for i in range(0, len(ids), 20):
        st, lote = ml.get("/items", ids=",".join(ids[i:i + 20]),
                          attributes="id,title,status,shipping,inventory_id,variations,seller_custom_field,date_created")
        for e in lote or []:
            if e.get("code") == 200:
                itens[e["body"]["id"]] = e["body"]
    return itens


def anuncios_full(ml: ML, seller_id: str) -> dict:
    ids, offset = [], 0
    while True:
        st, b = ml.get(f"/users/{seller_id}/items/search", status="active", limit=100, offset=offset)
        res = (b or {}).get("results", [])
        ids += res
        offset += len(res)
        if not res or offset >= b["paging"]["total"] or offset >= 1000:
            break
    itens = detalhar(ml, ids)
    return {k: v for k, v in itens.items() if (v.get("shipping") or {}).get("logistic_type") == "fulfillment"}


def mapa_inventario(itens: dict) -> dict:
    """(item_id, variation_id|None) → inventory_id"""
    m = {}
    for it in itens.values():
        if it.get("inventory_id"):
            m[(it["id"], None)] = it["inventory_id"]
        for v in it.get("variations") or []:
            if v.get("inventory_id"):
                m[(it["id"], v.get("id"))] = v["inventory_id"]
    return m


def nome_variacao(var: dict) -> str:
    """'Cor da caixa: Preto · Cor da pulseira: Preto'"""
    return " · ".join(f"{a.get('name')}: {a.get('value_name')}" for a in var.get("attribute_combinations") or [])


def vendas_por_item(ml: ML, seller_id: str, ate: datetime.date):
    """Retorna ({(item_id, variation_id): {dia: unidades}}, {(item_id, variation_id): seller_sku})
    — 30 dias até `ate`, consultando dia a dia."""
    out = defaultdict(lambda: defaultdict(int))
    skus = {}
    for d in range(30):
        dia = ate - timedelta(days=d)
        offset = 0
        while True:
            st, b = ml.get("/orders/search", seller=seller_id, limit=50, offset=offset, sort="date_asc",
                           **{"order.date_created.from": f"{dia}T00:00:00.000-04:00",
                              "order.date_created.to": f"{dia}T23:59:59.000-04:00"})
            if st != 200:
                print(f"   [{ml.conta}] orders {dia} offset {offset}: HTTP {st}")
                break
            res = b.get("results", [])
            for o in res:
                if o.get("status") in STATUS_FORA:
                    continue
                if (o.get("cancel_detail") or {}).get("code") == "pack_splitted":
                    continue
                ds = datetime.fromisoformat(o["date_created"]).astimezone(ML_TZ).date()
                for it in o.get("order_items", []):
                    item = it.get("item") or {}
                    chave = (item.get("id"), item.get("variation_id"))
                    out[chave][ds] += it.get("quantity", 0)
                    if item.get("seller_sku"):
                        skus[chave] = item["seller_sku"]
            offset += len(res)
            if not res or offset >= b["paging"]["total"]:
                break
    return out, skus


def dias_ruptura(ml: ML, seller_id: str, inv: str, ate) -> set:
    """Dias (fuso -04:00) em que o inventário ficou o dia INTEIRO sem unidade disponível.
    Reconstrói o saldo diário de available_quantity pelas operações do Full (30 dias)."""
    de = ate - timedelta(days=29)
    ops, scroll = [], None
    while True:
        params = {"seller_id": seller_id, "inventory_id": inv,
                  "date_from": de.isoformat(), "date_to": (ate + timedelta(days=1)).isoformat()}
        if scroll:
            params["scroll"] = scroll
        st, b = ml.get("/stock/fulfillment/operations/search", **params)
        if st != 200 or not b:
            return None  # sem dado confiável (ex.: limite da API): não ajusta e sinaliza
        res = b.get("results", [])
        ops += res
        scroll = (b.get("paging") or {}).get("scroll")
        if not res or not scroll or len(ops) >= (b.get("paging") or {}).get("total", 0):
            break
    if not ops:
        return SEM_MOVIMENTO  # nenhuma operação no Full em 30 dias
    ops.sort(key=lambda o: o["date_created"])
    # saldo antes da 1ª operação = resultado − variação dela
    primeiro = ops[0]
    saldo = ((primeiro.get("result") or {}).get("available_quantity", 0)
             - (primeiro.get("detail") or {}).get("available_quantity", 0))
    por_dia = defaultdict(list)
    for o in ops:
        d = datetime.fromisoformat(o["date_created"].replace("Z", "+00:00")).astimezone(ML_TZ).date()
        por_dia[d].append((o.get("result") or {}).get("available_quantity", 0))
    zerados = set()
    for i in range(30):
        d = de + timedelta(days=i)
        valores = [saldo] + por_dia.get(d, [])  # saldo no início do dia + saldos após cada operação
        if max(valores) <= 0:
            zerados.add(d)
        if por_dia.get(d):
            saldo = por_dia[d][-1]
    return zerados


def estoque(ml: ML, inv: str) -> dict:
    st, b = ml.get(f"/inventories/{inv}/stock/fulfillment")
    if st != 200 or not b:
        return {"erro": st}
    det = {(x.get("status") or "").lower().replace("_", ""): x.get("quantity", 0) for x in b.get("not_available_detail") or []}
    return {"disponivel": b.get("available_quantity", 0), "transferencia": det.get("transfer", 0),
            "outros_indisp": (b.get("not_available_quantity", 0) or 0) - det.get("transfer", 0), "total": b.get("total", 0)}


def calcular(conta: str, ate) -> list:
    cfg = ml_api.load_config()["contas"][conta]
    seller_id = str(cfg["seller_id"])
    ml = ML(conta)
    t0 = time.time()
    full = anuncios_full(ml, seller_id)
    vendas, skus = vendas_por_item(ml, seller_id, ate)
    # MLBs vendidos que não estão entre os Full ativos (ex.: pausados que dividem inventário)
    faltam = sorted({k[0] for k in vendas if k[0] and k[0] not in full})
    extras = {k: v for k, v in detalhar(ml, faltam).items() if v.get("inventory_id") or
              any(x.get("inventory_id") for x in v.get("variations") or [])}
    todos = {**full, **extras}
    mapa = mapa_inventario(todos)

    inv = defaultdict(lambda: {"titulo": "", "dias": defaultdict(int), "vinculos": {}})
    for (item_id, var_id), invid in mapa.items():
        it = todos[item_id]
        var = next((v for v in it.get("variations") or [] if v.get("id") == var_id), {})
        reg = inv[invid]
        reg["vinculos"][(item_id, var_id)] = {
            "mlb": item_id, "variacao_id": var_id, "variacao": nome_variacao(var),
            "status": it.get("status"), "full": item_id in full, "titulo": it.get("title", ""),
            "sku": var.get("seller_custom_field") or it.get("seller_custom_field") or skus.get((item_id, var_id)) or "",
            "un_30d": 0,
            "criado": (datetime.fromisoformat(it["date_created"].replace("Z", "+00:00")).astimezone(ML_TZ).date()
                       if it.get("date_created") else None),
        }

    # cada venda vai para o inventário da variação vendida; se não houver, para o do anúncio
    for (item_id, var_id), dias in vendas.items():
        chave = (item_id, var_id) if (item_id, var_id) in mapa else (item_id, None)
        invid = mapa.get(chave)
        if not invid:
            continue  # venda de anúncio fora do Full
        vinc = inv[invid]["vinculos"][chave]
        vinc["sku"] = vinc["sku"] or skus.get((item_id, var_id), "")
        for d, q in dias.items():
            inv[invid]["dias"][d] += q
            vinc["un_30d"] += q

    # MLB principal = anúncio Full ativo que mais vendeu esse estoque
    for reg in inv.values():
        vs = sorted(reg["vinculos"].values(), key=lambda v: (not (v["full"] and v["status"] == "active"), -v["un_30d"]))
        reg["principal"] = vs[0] if vs else {}
        reg["mlbs"] = {v["mlb"] for v in vs}
        reg["detalhe_mlbs"] = " | ".join(
            f"{v['mlb']} ({'Full ' if v['full'] else ''}{STATUS_PT.get(v['status'], v['status'])}, {v['un_30d']} un/30d)"
            for v in vs)

    linhas = []
    for invid, reg in inv.items():
        if not any(m in full for m in reg["mlbs"]):
            continue  # inventário sem nenhum anúncio Full ativo
        est = estoque(ml, invid)
        un = {j: sum(q for d, q in reg["dias"].items() if d > ate - timedelta(days=j)) for j in (7, 15, 30)}

        # Ponderação 1 — produto novo: dias antes do anúncio mais antigo do inventário não contam
        criados = [v["criado"] for v in reg["vinculos"].values() if v.get("criado")]
        desde = min(criados) if criados else None
        # Ponderação 2 — ruptura: dias inteiros sem estoque disponível não contam.
        # Só consulta as operações (rate limit) para inventários suspeitos.
        suspeito = un[30] > 0 and ((est.get("disponivel") or 0) == 0 or un[7] < 0.5 * 7 * un[30] / 30)
        zerados = dias_ruptura(ml, seller_id, invid, ate) if suspeito else set()
        nao_verificado = zerados is None
        sem_movimento = zerados == SEM_MOVIMENTO
        zerados = zerados if isinstance(zerados, set) else set()
        validos = {}
        for j in (7, 15, 30):
            dias_j = [ate - timedelta(days=i) for i in range(j)]
            validos[j] = sum(1 for d in dias_j if (desde is None or d >= desde) and d not in zerados)
        med = {j: un[j] / max(validos[j], 1) for j in un}
        # janela só vale com pelo menos metade dos dias válidos (evita média de 1 ou 2 dias)
        confiavel = {j: validos[j] >= math.ceil(j / 2) for j in un}
        if not any(confiavel.values()):
            confiavel[max(validos, key=validos.get)] = True  # usa a janela com mais dias válidos
        obs = []
        if nao_verificado:
            obs.append("ruptura não verificada (limite da API)")
        if sem_movimento and (est.get("disponivel") or 0) + (est.get("transferencia") or 0) == 0:
            obs.append("sem estoque e sem movimento no Full há 30 dias — vendas vieram de outra logística")
        if desde and desde > ate - timedelta(days=30):
            obs.append(f"produto novo ({(ate - desde).days + 1} dias)")
        if zerados:
            obs.append(f"ruptura ajustada ({len(zerados)} dias zerado)")
        if confiavel[7] and confiavel[30]:
            tend = med[7] / med[30] if med[30] > 0 else (math.inf if med[7] > 0 else 1.0)
        else:
            tend = 1.0  # sem base para medir tendência
        pesos = {j: p for j, p in PESOS.items() if confiavel[j]}
        demanda = sum(med[j] * p for j, p in pesos.items()) / sum(pesos.values())
        if tend > ALTA:
            demanda = max(demanda, *(med[j] for j in pesos))
        descartadas = [f"{j}d" for j in PESOS if not confiavel[j]]
        if descartadas:
            obs.append(f"janela {', '.join(descartadas)} ignorada (poucos dias válidos)")
        considerado = (est.get("disponivel") or 0) + (est.get("transferencia") or 0)
        cobertura = considerado / demanda if demanda > 0 else math.inf
        sugestao = max(0, math.ceil(META_DIAS * demanda - considerado))
        p = reg["principal"]
        linhas.append({
            "conta": conta, "inventory_id": invid, "mlb": p.get("mlb", ""), "variacao": p.get("variacao", ""),
            "sku": p.get("sku", ""), "titulo": p.get("titulo", ""), "mlbs": reg["detalhe_mlbs"],
            "n_mlbs": len(reg["mlbs"]), "un_7d": un[7], "un_15d": un[15], "un_30d": un[30],
            "media_7d": round(med[7], 2), "media_15d": round(med[15], 2), "media_30d": round(med[30], 2),
            "tendencia": "alta" if tend > ALTA else "queda" if tend < QUEDA else "estável",
            "tend_7x30": None if math.isinf(tend) else round(tend, 2),
            "demanda_dia": round(demanda, 2), "disponivel": est.get("disponivel"),
            "transferencia": est.get("transferencia"), "outros_indisp": est.get("outros_indisp"),
            "estoque_considerado": considerado,
            "cobertura_dias": None if math.isinf(cobertura) else round(cobertura, 1),
            "sugestao_reposicao": sugestao, "erro_estoque": est.get("erro"),
            "anuncio_desde": desde.strftime("%d/%m/%Y") if desde else "", "dias_zerado": len(zerados),
            "validos": f"{validos[7]} / {validos[15]} / {validos[30]}", "obs": "; ".join(obs),
        })
    print(f"{conta}: {len(full)} anúncios Full, {len(linhas)} inventários, "
          f"{sum(1 for l in linhas if l['sugestao_reposicao'] > 0)} com reposição sugerida ({time.time() - t0:.0f}s)")
    return linhas


def salvar_xlsx(linhas: list, caminho: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Reposição Full"
    cab = ["Conta", "MLB principal", "Variação", "SKU", "Produto", "Inventário", "Qtd. MLBs",
           "MLBs vinculados (status, vendas 30d)", "Un. 7d", "Un. 15d", "Un. 30d",
           "Média/dia 7d", "Média/dia 15d", "Média/dia 30d", "Tendência", "7d ÷ 30d", "Demanda/dia",
           "Disponível", "Em transferência", "Outros indisp.", "Estoque considerado",
           "Cobertura (dias)", f"Sugestão p/ {META_DIAS} dias",
           "Anúncio desde", "Dias zerado (30d)", "Dias válidos 7/15/30", "Observação"]
    chaves = ["conta", "mlb", "variacao", "sku", "titulo", "inventory_id", "n_mlbs", "mlbs",
              "un_7d", "un_15d", "un_30d",
              "media_7d", "media_15d", "media_30d", "tendencia", "tend_7x30", "demanda_dia",
              "disponivel", "transferencia", "outros_indisp", "estoque_considerado",
              "cobertura_dias", "sugestao_reposicao",
              "anuncio_desde", "dias_zerado", "validos", "obs"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    verm, amar = PatternFill("solid", fgColor="FCEBEB"), PatternFill("solid", fgColor="FDF4DC")
    for l in linhas:
        ws.append([l[k] for k in chaves])
        if l["mlb"]:
            cel = ws.cell(row=ws.max_row, column=2)
            cel.hyperlink = f"https://produto.mercadolivre.com.br/MLB-{l['mlb'][3:]}"
            cel.style = "Hyperlink"
        cob = l["cobertura_dias"]
        if cob is not None and cob < META_DIAS:
            for c in ws[ws.max_row]:
                c.fill = verm if cob < META_DIAS / 2 else amar
    larg = {"Produto": 50, "Variação": 30, "MLBs vinculados (status, vendas 30d)": 60, "MLB principal": 16,
            "Observação": 40, "Dias válidos 7/15/30": 18}
    for i, h in enumerate(cab, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = larg.get(h, 14)
    ws.freeze_panes = "E2"
    wb.save(caminho)


if __name__ == "__main__":
    contas = sys.argv[1].split(",") if len(sys.argv) > 1 else ["meli01", "meli02", "meli03"]
    ate = datetime.now(ML_TZ).date() - timedelta(days=1)  # último dia fechado
    print(f"Janelas terminando em {ate} (fuso -04:00) | meta {META_DIAS} dias\n")
    linhas = []
    for conta in contas:
        linhas += calcular(conta, ate)
    ordem = lambda l: (l["cobertura_dias"] if l["cobertura_dias"] is not None else 1e9)
    linhas.sort(key=ordem)
    caminho = os.path.join(BASE_DIR, f"reposicao_full_{ate:%Y%m%d}.xlsx")
    salvar_xlsx(linhas, caminho)

    print(f"\n{'conta':6} {'MLB':14} {'produto':30} {'variação':22} {'7d':>4} {'15d':>4} {'30d':>4} {'tend.':>7} "
          f"{'dem/d':>6} {'disp':>5} {'transf':>6} {'cob(d)':>7} {'repor':>6}")
    for l in linhas:
        if l["sugestao_reposicao"] <= 0:
            continue
        cob = "—" if l["cobertura_dias"] is None else f"{l['cobertura_dias']:.1f}"
        mlb = l["mlb"] + ("+" if l["n_mlbs"] > 1 else "")
        print(f"{l['conta']:6} {mlb:14} {l['titulo'][:30]:30} {(l['variacao'] or '—')[:22]:22} {l['un_7d']:>4} {l['un_15d']:>4} "
              f"{l['un_30d']:>4} {l['tendencia']:>7} {l['demanda_dia']:>6.2f} {l['disponivel'] or 0:>5} "
              f"{l['transferencia'] or 0:>6} {cob:>7} {l['sugestao_reposicao']:>6}  {l['obs']}")
    print(f"\nPlanilha completa (todos os inventários): {caminho}")
